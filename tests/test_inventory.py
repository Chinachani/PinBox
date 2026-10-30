import copy
import io
import json
import sqlite3
import hashlib
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import openpyxl
import pytest

import database
import main
from bom import decode_spec, parse_bom_text
from parser import parse_lcsc_barcode


def create(client, **fields):
    payload = {"code": "C10001", "model": "TEST-123", "quantity": 20, **fields}
    response = client.post("/api/materials", json=payload)
    assert response.status_code == 200, response.text
    return payload


def snapshot(client):
    return client.get("/api/backup").json()


def restore(client, data):
    return client.post("/api/restore?confirm=yes", files={"file": ("backup.json", json.dumps(data).encode())})


def push(client, operation, dataset=None):
    dataset = dataset or client.get("/api/sync/state").json()["dataset_id"]
    return client.post("/api/sync/op", json={"operation": operation, "dataset_id": dataset})


@pytest.mark.parametrize("field,value", [("quantity", -1), ("quantity", 1.5), ("quantity", True), ("min_stock", -1), ("price", -0.1), ("model", "  "), ("code", "  ")])
def test_invalid_create_does_not_write(client, field, value):
    response = client.post("/api/materials", json={"code": "C10001", "model": "test", field: value})
    assert response.status_code == 422
    assert client.get("/api/stats").json()["total_sku"] == 0


def test_quantity_update_and_pagination_are_validated(client):
    create(client)
    for payload in ({"quantity": -1}, {"quantity": 1.2}, {"price": -1}, {"model": " "}):
        assert client.put("/api/materials/C10001", json=payload).status_code == 422
    for path in ("/api/materials?page_size=0", "/api/logs?page_size=-1", "/api/materials?page=0"):
        assert client.get(path).status_code == 422
    assert client.get("/api/materials/C10001").json()["quantity"] == 20


def test_concurrent_stock_never_overdraws_and_logs_match(client):
    create(client)
    with ThreadPoolExecutor(max_workers=8) as pool:
        responses = list(pool.map(lambda _: client.post("/api/materials/C10001/stock", json={"delta": -1}), range(30)))
    assert sum(r.status_code == 200 for r in responses) == 20
    assert all(r.status_code in (200, 400) for r in responses)
    successes = [r.json() for r in responses if r.status_code == 200]
    assert sorted(r["previous"] for r in successes) == list(range(1, 21))
    material = client.get("/api/materials/C10001").json()
    assert material["quantity"] == 0
    assert len(material["logs"]) == 21
    assert material["logs"][0]["remaining"] == 0


def test_failed_log_insert_rolls_back_stock(client):
    create(client)
    with database.transaction() as conn:
        conn.execute("CREATE TRIGGER fail_log BEFORE INSERT ON stock_logs BEGIN SELECT RAISE(ABORT, 'disk error'); END")
    assert client.post("/api/materials/C10001/stock", json={"delta": -1}).status_code == 500
    assert client.get("/api/materials/C10001").json()["quantity"] == 20


def test_sync_retry_is_exactly_once_even_concurrently(client):
    create(client)
    dataset = client.get("/api/sync/state").json()["dataset_id"]
    op = {"op_id": str(uuid4()), "kind": "stock", "code": "C10001", "delta": -3, "project": "board-A"}
    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(lambda _: push(client, op, dataset), range(4)))
    assert all(r.status_code == 200 for r in responses)
    material = client.get("/api/materials/C10001").json()
    assert material["quantity"] == 17
    assert len(material["logs"]) == 2
    assert material["logs"][0]["project"] == "board-A"
    assert push(client, {**op, "delta": -4}, dataset).status_code == 409


def test_sync_conflicts_never_overwrite_stock_or_metadata(client):
    create(client)
    op = {"op_id": str(uuid4()), "kind": "create", "payload": {"code": "C10001", "model": "offline", "quantity": 1}}
    assert push(client, op).status_code == 409
    update = {"op_id": str(uuid4()), "kind": "update", "code": "C10001", "payload": {"model": "offline"}, "expected": {"model": "old"}}
    assert push(client, update).status_code == 409
    assert client.get("/api/materials/C10001").json()["quantity"] == 20
    assert client.get("/api/materials/C10001").json()["model"] == "TEST-123"


def test_sync_resolves_location_names_not_local_ids(client):
    create(client, location="A")
    client.post("/api/locations", json={"name": "B"})
    op = {"op_id": str(uuid4()), "kind": "location_update", "id": 2, "original_name": "A", "expected_note": "", "payload": {"name": "A-renamed"}}
    assert push(client, op).status_code == 200
    assert client.get("/api/materials/C10001").json()["location"] == "A-renamed"
    assert any(loc["name"] == "B" for loc in client.get("/api/locations").json())


def test_sync_rejects_old_dataset_after_wipe(client):
    create(client)
    dataset = client.get("/api/sync/state").json()["dataset_id"]
    assert client.post("/api/wipe?confirm=yes").status_code == 200
    create(client)
    assert push(client, {"op_id": str(uuid4()), "kind": "stock", "code": "C10001", "delta": -1}, dataset).status_code == 409


def test_backup_roundtrip_preserves_project_ids_and_empty_database(client):
    create(client, location="A", price=1.25)
    client.post("/api/materials/C10001/stock", json={"delta": -2, "project": "board-A"})
    before = snapshot(client)
    assert restore(client, before).status_code == 200
    after = snapshot(client)
    for key in ("materials", "stock_logs", "locations"):
        assert after[key] == before[key]
    assert restore(client, {"materials": [], "locations": [{"id": 9, "name": "empty"}]}).status_code == 200
    assert client.get("/api/stats").json()["total_sku"] == 0


@pytest.mark.parametrize("corrupt", ["duplicate", "orphan", "negative", "bad_receipts", "wrong_structure"])
def test_invalid_restore_keeps_original_database(client, corrupt):
    create(client)
    before = snapshot(client)
    data = copy.deepcopy(before)
    if corrupt == "duplicate":
        data["materials"].append(data["materials"][0])
    elif corrupt == "orphan":
        data["stock_logs"][0]["material_code"] = "missing"
    elif corrupt == "negative":
        data["materials"][0]["quantity"] = -10
    elif corrupt == "bad_receipts":
        data["import_receipts"] = 123
    else:
        data = []
    assert restore(client, data).status_code == 400
    assert snapshot(client)["materials"] == before["materials"]


def test_csv_import_roundtrip_and_duplicate_detection(client):
    original = create(client, category="resistor", brand="brand", package="0603", quantity=0, price=0.125, min_stock=12, remark='hello,"quoted"\nnext line', location="A", use_case="test", order_date="2026-10-20")
    content = client.get("/api/export?format=csv").content
    assert client.post("/api/wipe?confirm=yes").status_code == 200
    response = client.post("/api/import", files={"file": ("stock.csv", content)})
    assert response.status_code == 200, response.text
    material = client.get("/api/materials/C10001").json()
    for key in original:
        assert material[key] == original[key], key
    assert client.post("/api/import", files={"file": ("renamed.csv", content)}).status_code == 409


def test_invalid_import_is_atomic(client):
    data = [{"code": "C10001", "model": "ok", "quantity": 10}, {"code": "C10002", "model": "bad", "quantity": -2}]
    response = client.post("/api/import", files={"file": ("stock.json", json.dumps(data).encode())})
    assert response.status_code == 400
    assert client.get("/api/stats").json()["total_sku"] == 0


def test_import_merges_batches_and_both_inventory_schemas(client):
    data = {"today_orders": [{"商品编号": "C10001", "商品型号/名称": "test", "采购数量": 2, "库存批次": "batch-A"}], "historical_inventory": [{"code": "C10001", "model": "test", "quantity": 3, "source": "batch-B"}]}
    response = client.post("/api/import", files={"file": ("stock.json", json.dumps(data).encode())})
    assert response.status_code == 200, response.text
    material = client.get("/api/materials/C10001").json()
    assert material["quantity"] == 5
    assert "batch-A" in material["batch"] and "batch-B" in material["batch"]
    assert len(material["logs"]) == 2


def test_import_receipt_deduplicates_server_and_multiple_offline_devices(client):
    records = [{"code": "C10001", "model": "test", "quantity": 5}]
    content = json.dumps(records).encode()
    digest = hashlib.sha256(content).hexdigest()
    assert client.post("/api/import", files={"file": ("orders.json", content)}).status_code == 200
    operation = {"op_id": str(uuid4()), "kind": "import", "records": records, "receipt": digest, "filename": "orders.json"}
    assert push(client, operation).status_code == 200
    assert push(client, {**operation, "op_id": str(uuid4())}).status_code == 200
    assert client.get("/api/materials/C10001").json()["quantity"] == 5


@pytest.mark.parametrize("text,code,qty", [('pc:C91372, pm:ESD-15, qty:10', 'C91372', 10), ('{"code":"C123","qty":5,"model":"test"}', 'C123', 5), ('C91372 40个', 'C91372', 40), ('C91372,100,20261026', 'C91372', 100), ('C91372,qty:1.5', 'C91372', None), ('C91372 1.5个', 'C91372', None), ('C91372 -2PCS', 'C91372', None), ('C91372 qty:1e3', 'C91372', None), ('C91372 qty:-2个', 'C91372', None), ('C91372 9999999999PCS', 'C91372', None), ('C91372 0PCS', 'C91372', None), ('C91372 数量:12个', 'C91372', 12)])
def test_barcode_formats(text, code, qty):
    result = parse_lcsc_barcode(text)
    assert result["code"] == code
    assert result["quantity"] == qty


def test_bom_quantity_specs_and_shared_stock(client):
    assert parse_bom_text('10k 0603 10000')[0]["qty"] == 10000
    assert parse_bom_text('10k 0603 10')[0]["model"] == '10k 0603'
    assert '4.7kΩ' in decode_spec('4K7')["res_values"]
    assert '1uF' in decode_spec('1µF')["cap_values"]
    create(client, quantity=10)
    response = client.post("/api/bom/check", json={"content": "C10001 6\nC10001 6"})
    assert response.status_code == 200, response.text
    items = response.json()["items"]
    assert [i["shortfall"] for i in items] == [0, 2]


def test_bom_suggestions_require_confirmation(client):
    create(client, model="10kΩ 0603 1%", category="电阻", quantity=100)
    response = client.post("/api/bom/check", json={"content": "10k 0603 10"})
    item = response.json()["items"][0]
    assert item["candidates"]
    assert item["matched_code"] is None


def test_bom_short_models_and_package_conflicts_require_confirmation(client):
    from bom import match_bom_items
    create(client, model='SS34 diode', package='SMA')
    item = match_bom_items([{'model': 'SS34', 'qty': 1, 'package': ''}])[0]
    assert item['matched_code'] is None
    assert item['candidates']
    client.put('/api/materials/C10001', json={'model': 'SS34'})
    assert match_bom_items([{'model': 'SS34', 'qty': 1, 'package': 'SMB'}])[0]['matched_code'] is None
    assert match_bom_items([{'model': 'SS34', 'qty': 1, 'package': 'SMA'}])[0]['matched_code'] == 'C10001'
    create(client, code='C10002', model='SS34', package='SMB')
    ambiguous = match_bom_items([{'model': 'SS34', 'qty': 1, 'package': ''}])[0]
    assert ambiguous['matched_code'] is None
    assert len(ambiguous['candidates']) == 2


def test_order_excel_columns_are_aligned(client):
    response = client.post("/api/bom/order-excel", json={"items": [{"status": "insufficient", "matched_code": "C10001", "matched_model": "TEST", "matched_brand": "Brand", "qty": 8, "shortfall": 3, "matched_price": 1.25}]})
    assert response.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert [wb.active.cell(3, i).value for i in range(1, 8)] == [1, "C10001", "TEST", "Brand", 3, 1.25, 3.75]


def test_excel_exports_user_text_as_text_instead_of_formulas(client):
    create(client, model='=2+2')
    response = client.get('/api/export/excel')
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert wb.active['D3'].value == '=2+2'
    assert wb.active['D3'].data_type == 's'


def test_initialization_refuses_silent_reset_and_backs_up_explicit_reset(client, tmp_path):
    import importlib.util
    from pathlib import Path
    spec = importlib.util.spec_from_file_location('pinbox_init', Path(__file__).resolve().parents[1] / 'init_data.py')
    initializer = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(initializer)
    create(client)
    source = tmp_path / 'source.json'
    source.write_text(json.dumps([{'code': 'C10002', 'model': 'replacement', 'quantity': 3}]))
    with pytest.raises(ValueError):
        initializer.load_initial_data(source)
    assert client.get('/api/materials/C10001').status_code == 200
    initializer.load_initial_data(source, reset=True)
    assert client.get('/api/materials/C10002').json()['quantity'] == 3
    backups = list((tmp_path / 'backups').glob('*.db'))
    assert backups
    with sqlite3.connect(backups[0]) as conn:
        assert conn.execute("SELECT quantity FROM materials WHERE code='C10001'").fetchone()[0] == 20


def test_api_token_and_cross_origin_write_protection(client, monkeypatch):
    monkeypatch.setenv("PINBOX_API_TOKEN", "test-token")
    assert client.get("/api/stats").status_code == 401
    assert client.get("/api/stats", headers={"Authorization": "Bearer test-token"}).status_code == 200
    preflight = client.options("/api/stats", headers={"Origin": "https://localhost", "Access-Control-Request-Method": "GET", "Access-Control-Request-Headers": "Authorization"})
    assert preflight.status_code == 200
    monkeypatch.delenv("PINBOX_API_TOKEN")
    assert client.post("/api/wipe?confirm=yes", headers={"Origin": "https://unrelated.example"}).status_code == 403


def test_live_notifications_cross_client_event_loops(client):
    from fastapi.testclient import TestClient
    create(client)
    with client.websocket_connect('/ws', headers={'origin': 'http://testserver'}) as first:
        with TestClient(main.app) as other:
            with other.websocket_connect('/ws', headers={'origin': 'http://testserver'}) as second:
                assert client.post('/api/materials/C10001/stock', json={'delta': -1}).status_code == 200
                assert first.receive_json()['current'] == 19
                assert second.receive_json()['current'] == 19


def test_windows_gui_bootloader_without_console_can_initialize_logging(tmp_path):
    import os
    import subprocess
    import sys
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    code = '''
import runpy, sys
sys.frozen = True
sys.stdout = sys.stderr = None
runpy.run_path('backend/main.py', run_name='pinbox_gui_test')
import uvicorn
config = uvicorn.Config('main:app')
assert sys.stdout is not None and sys.stderr is not None
'''
    subprocess.run([sys.executable, '-c', code], cwd=root, check=True, timeout=20,
                   env={**os.environ, 'PINBOX_DATA_DIR': str(tmp_path)})
